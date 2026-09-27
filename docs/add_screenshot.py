# -*- coding: utf-8 -*-
"""Adds a real screenshot of the completed, running mobile application to
Chapter 5, right after the Functional Testing table, as genuine evidence of
the implemented system (as opposed to the Chapter 4 design-stage wireframes).
"""
import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

PATH = "../159056-DOCUMENTATION.docx"
doc = docx.Document(PATH)

# Anchor: the "Results and Discussion" Heading 2 in Chapter 5 - insert the
# figure immediately before it, right after the functional testing table.
anchor = None
for p in doc.paragraphs:
    if p.style.name == "Heading 2" and p.text.strip() == "Results and Discussion":
        anchor = p
        break
assert anchor is not None, "Could not find Results and Discussion anchor"

img_p = anchor.insert_paragraph_before()
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = img_p.add_run()
run.add_picture("app_screenshot.png", width=Inches(2.6))

cap_p = anchor.insert_paragraph_before(
    "Figure 5.1: Sign-In Screen of the Completed UniGuide Mobile Application, Running on an Android Emulator",
    style="Caption",
)
cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

spacer = anchor.insert_paragraph_before("")

doc.save(PATH)
print("Saved:", PATH)
