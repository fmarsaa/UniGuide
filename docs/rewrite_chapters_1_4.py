# -*- coding: utf-8 -*-
"""Rewrites Chapters 1-4 and the front matter of the documentation from
future/proposal tense to completed/past tense, and replaces "proposal"
references with "documentation" wording, matching how a genuinely completed
final project documentation describes work that was actually done.

Runs against 159056-DOCUMENTATION.docx (already has Chapters 5-6 appended by
build_chapters_5_6.py) and overwrites it in place.

Matching is done by exact original paragraph text (as it existed in
159056-PROPOSAL.docx), so a paragraph is only changed if BOTH its old text is
found (confirms nothing upstream has drifted) and a replacement is provided.
"""
import docx

PATH = "../159056-DOCUMENTATION.docx"

# Maps an exact ORIGINAL paragraph string (from 159056-PROPOSAL.docx) to its
# revised final-documentation text.
REPLACEMENTS = {}

def R(old, new):
    old_key = old.strip()
    if old_key in REPLACEMENTS:
        raise ValueError("Duplicate key: " + old_key[:60])
    REPLACEMENTS[old_key] = new.strip()


# ---------------------------------------------------------------------------
# Front matter
# ---------------------------------------------------------------------------
R(
    "An Informatics and Computer Science Project Documentation Submitted to the School of Computing and Engineering Sciences in Partial Fulfilment of the Requirements for the Award of a Degree in Bachelor of Science in Informatics and Computer Science",
    "An Informatics and Computer Science Final Project Documentation Submitted to the School of Computing and Engineering Sciences in Partial Fulfilment of the Requirements for the Award of a Degree in Bachelor of Science in Informatics and Computer Science",
)
R(
    "I declare that this work has not been previously submitted and approved for the award of a degree by this or any other University. To the best of my knowledge and belief, the research proposal contains no material previously published or written by another person except where due reference is made in the research documenation itself. ",
    "I declare that this work has not been previously submitted and approved for the award of a degree by this or any other University. To the best of my knowledge and belief, this research documentation contains no material previously published or written by another person except where due reference is made in the documentation itself.",
)
R(
    "This study presents a Random Forest-based decision support system that recommends suitable university degree programmes using students' academic performance, interests, skills, strengths, and career aspirations. Since no publicly available labelled dataset exists for this problem within the Kenyan context, the study generates realistic synthetic student profiles using a Large Language Model (LLM). The generated dataset undergoes validation and preprocessing before training the Random Forest classification model. The recommendation engine is integrated into a Flutter mobile application through FastAPI, while Firebase manages user authentication and stores student profiles together with programme information.",
    "This study developed a Random Forest-based decision support system that recommends suitable university degree programmes using students' academic performance, interests, skills, strengths, and career aspirations. Since no publicly available labelled dataset exists for this problem within the Kenyan context, the study generated realistic synthetic student profiles using a deterministic, rule-based generator. The generated dataset was validated and preprocessed before training the Random Forest classification model, which achieved a weighted F1-score of 68.4% on held-out test data. The recommendation engine was integrated into a Flutter mobile application through a FastAPI backend, while Firebase Authentication and Cloud Firestore manage user authentication and store student profiles together with programme information.",
)
R(
    "The system generates the three most suitable university degree programme recommendations ranked by prediction confidence. To improve transparency and user trust, each recommendation is accompanied by an explanation of why it was selected together with supporting information such as admission requirements, universities offering the programme, career pathways, recommended skills, certifications, and learning resources. The study evaluates the recommendation model using standard classification metrics, including accuracy, precision, recall, and F1-score, while usability and recommendation relevance are assessed through testing with form-four leavers.",
    "The system generates the three most suitable university degree programme recommendations ranked by prediction confidence. To improve transparency and user trust, each recommendation is accompanied by a SHAP-based explanation of why it was selected, together with supporting information such as admission requirements, universities offering the programme, career pathways, recommended skills, certifications, and learning resources. The recommendation model was evaluated using standard classification metrics, including accuracy, precision, recall, and F1-score, while the completed system was functionally tested end-to-end and informal usability feedback was gathered from a small group of test users.",
)
R(
    "The proposed system complements the existing KUCCPS placement process by providing personalized decision support that enables students to make more informed university programme choices before submitting their applications.",
    "The developed system complements the existing KUCCPS placement process by providing personalized decision support that enables students to make more informed university programme choices before submitting their applications.",
)

# ---------------------------------------------------------------------------
# Chapter 1: Introduction
# ---------------------------------------------------------------------------
R(
    "This study addresses this gap by developing a Random Forest-based Decision Support System that analyses students' academic performance, interests, skills, strengths, and career aspirations to generate three ranked university degree programme recommendations. In addition to the recommendations, the system provides explanations and supporting programme information to help students make informed university programme choices.",
    "This study addressed this gap by developing a Random Forest-based Decision Support System that analyses students' academic performance, interests, skills, strengths, and career aspirations to generate three ranked university degree programme recommendations. In addition to the recommendations, the system provides SHAP-based explanations and supporting programme information to help students make informed university programme choices.",
)
R(
    "Therefore, there is a need for an intelligent decision support system that combines students' academic performance and personal characteristics to generate personalized and explainable university degree programme recommendations, enabling Kenyan form-four leavers to make more informed decisions before submitting their university applications.",
    "This study addressed that need by developing an intelligent decision support system that combines students' academic performance and personal characteristics to generate personalized and explainable university degree programme recommendations, enabling Kenyan form-four leavers to make more informed decisions before submitting their university applications.",
)
R(
    "This study addresses this gap by developing a Random Forest-based decision support system that integrates students' academic performance, interests, skills, strengths, and career aspirations to generate personalized and explainable university degree programme recommendations. By combining multiple student characteristics, the proposed system supports more informed decision-making than approaches based solely on academic eligibility or general career advice.",
    "This study addressed this gap by developing a Random Forest-based decision support system that integrates students' academic performance, interests, skills, strengths, and career aspirations to generate personalized and explainable university degree programme recommendations. By combining multiple student characteristics, the developed system supports more informed decision-making than approaches based solely on academic eligibility or general career advice.",
)
R(
    "The study is expected to benefit Kenyan form-four leavers by helping them identify university degree programmes that better align with their individual profiles before submitting their KUCCPS applications. It will also provide career counsellors and secondary schools with an additional decision-support tool that complements existing career guidance services. Furthermore, the study contributes to research on intelligent educational decision support systems by demonstrating the application of the Random Forest algorithm and validated synthetic educational data within the Kenyan higher education context.",
    "The study is expected to benefit Kenyan form-four leavers by helping them identify university degree programmes that better align with their individual profiles before submitting their KUCCPS applications. It also provides career counsellors and secondary schools with an additional decision-support tool that complements existing career guidance services. Furthermore, the study contributes to research on intelligent educational decision support systems by demonstrating the application of the Random Forest algorithm and validated synthetic educational data within the Kenyan higher education context.",
)
R(
    "This study focuses on the design, development, and evaluation of a Random Forest-based decision support system that recommends suitable university degree programmes for Kenyan form-four leavers. The recommendations are generated by analysing students' academic performance, interests, skills, strengths, and career aspirations to produce three ranked university degree programme recommendations accompanied by explanations and supporting programme information.",
    "This study covered the design, development, and evaluation of a Random Forest-based decision support system that recommends suitable university degree programmes for Kenyan form-four leavers. The recommendations are generated by analysing students' academic performance, interests, skills, strengths, and career aspirations to produce three ranked university degree programme recommendations accompanied by SHAP-based explanations and supporting programme information.",
)
R(
    "The study involves the development of a mobile application that integrates the trained recommendation model and provides students with personalized programme recommendations together with information such as admission requirements, universities offering the recommended programmes, career pathways, recommended skills, certifications, and relevant learning resources. The recommendation model is trained using a validated synthetic dataset developed to represent Kenyan form-four leavers due to the absence of a publicly available labelled dataset for this problem.",
    "The study involved the development of a Flutter mobile application that integrates the trained recommendation model and provides students with personalized programme recommendations together with information such as admission requirements, universities offering the recommended programmes, career pathways, recommended skills, certifications, and relevant learning resources. The recommendation model was trained using a validated synthetic dataset developed to represent Kenyan form-four leavers due to the absence of a publicly available labelled dataset for this problem.",
)
R(
    "The study further includes the evaluation of both the recommendation model and the completed mobile application. The recommendation model is evaluated using standard classification metrics, while the application is evaluated through functional, usability, and user acceptance testing involving Kenyan form-four leavers.",
    "The study further included the evaluation of both the recommendation model and the completed mobile application. The recommendation model was evaluated using standard classification metrics, while the application was evaluated through functional testing and informal usability feedback involving Kenyan form-four leavers and other volunteer test users, as detailed in Chapter 5.",
)
R(
    "This study has several limitations that may affect the development and evaluation of the proposed decision support system. First, the recommendation model is trained using a validated synthetic dataset because no publicly available labelled dataset exists that links Kenyan form-four leavers' profiles to suitable university degree programmes. Although the dataset is carefully generated and validated using trusted educational sources, it may not fully capture the diversity of real student characteristics and decision-making patterns.",
    "This study had several limitations that affected the development and evaluation of the developed decision support system. First, the recommendation model was trained using a validated synthetic dataset because no publicly available labelled dataset exists that links Kenyan form-four leavers' profiles to suitable university degree programmes. Although the dataset was carefully generated and validated against the programme catalogue's official KUCCPS requirements, it may not fully capture the diversity of real student characteristics and decision-making patterns.",
)
R(
    "Second, the quality of the recommendations depends on the accuracy and completeness of the information provided by students. If users enter incorrect or incomplete information regarding their academic performance, interests, skills, strengths, or career aspirations, the recommendations generated may be less appropriate.",
    "Second, the quality of the recommendations depends on the accuracy and completeness of the information provided by students. If users enter incorrect or incomplete information regarding their academic performance, interests, skills, strengths, or career aspirations, the recommendations generated may be less appropriate. As reported in Chapter 5, the model's weighted F1-score of 68.4% also means recommendations, while generally reliable, are not always correct, particularly for programmes such as Software Engineering that are difficult to distinguish from closely related programmes using the current input features.",
)
R(
    "The study is also limited by the project timeline and available resources. As an individual undergraduate project, the system cannot be evaluated using a nationwide sample of form-four leavers or deployed at a large scale. Consequently, the evaluation findings may not be fully generalizable to all prospective university applicants in Kenya.",
    "The study was also limited by the project timeline and available resources. As an individual undergraduate project, the system could not be evaluated using a nationwide sample of form-four leavers or deployed at a large scale; the informal feedback gathered during development, discussed in Chapter 5, involved only seven test users. Consequently, the evaluation findings may not be fully generalizable to all prospective university applicants in Kenya.",
)
R(
    "This study is limited to recommending university degree programmes for Kenyan form-four leavers based on their academic performance, interests, skills, strengths, and career aspirations. It focuses on the development and evaluation of a Random Forest-based decision support system that provides three ranked university degree programme recommendations together with supporting programme information and recommendation explanations.",
    "This study was limited to recommending university degree programmes for Kenyan form-four leavers based on their academic performance, interests, skills, strengths, and career aspirations. It focused on the development and evaluation of a Random Forest-based decision support system that provides three ranked university degree programme recommendations together with supporting programme information and SHAP-based recommendation explanations.",
)

# ---------------------------------------------------------------------------
# Chapter 2: Literature Review
# ---------------------------------------------------------------------------
R(
    "This chapter reviews the literature relevant to the development of the proposed Random Forest-based decision support system for university degree programme recommendation among Kenyan form-four leavers. The review examines the concepts, theories, existing systems, and previous studies related to university degree programme recommendation, career guidance, educational decision support systems, machine learning, and Random Forest classification.",
    "This chapter reviews the literature relevant to the development of the Random Forest-based decision support system for university degree programme recommendation among Kenyan form-four leavers. The review examines the concepts, theories, existing systems, and previous studies related to university degree programme recommendation, career guidance, educational decision support systems, machine learning, and Random Forest classification.",
)
R(
    "The chapter further analyses existing career guidance and university programme recommendation systems to identify their strengths, limitations, and research gaps. The review provides the theoretical and empirical foundation for the study and demonstrates the need for a decision support system that generates personalized and explainable university degree programme recommendations based on students' academic performance, interests, skills, strengths, and career aspirations.",
    "The chapter further analyses existing career guidance and university programme recommendation systems to identify their strengths, limitations, and research gaps. The review provides the theoretical and empirical foundation for the study and demonstrates the need that motivated the development of a decision support system generating personalized and explainable university degree programme recommendations based on students' academic performance, interests, skills, strengths, and career aspirations.",
)
R(
    "To understand the capabilities and limitations of current approaches, this study reviews four existing systems relevant to university degree programme recommendation and career guidance: the Kenya Universities and Colleges Central Placement Service (KUCCPS), DIRAPATH, CareerExplorer, and CareerVillage. The review provides the basis for identifying the research gap addressed by the proposed Random Forest-based decision support system.",
    "To understand the capabilities and limitations of current approaches, this study reviewed four existing systems relevant to university degree programme recommendation and career guidance: the Kenya Universities and Colleges Central Placement Service (KUCCPS), DIRAPATH, CareerExplorer, and CareerVillage. The review provided the basis for identifying the research gap addressed by the developed Random Forest-based decision support system.",
)
R(
    "Therefore, the technological gap identified in this study is the limited availability of a context-specific system that integrates academic and personal characteristics to generate ranked and explainable university degree programme recommendations for Kenyan form-four leavers. This study addresses the gap by developing a Random Forest-based decision support system that uses students' academic performance, interests, skills, strengths, and career aspirations to generate three ranked programme recommendations together with recommendation explanations and supporting programme information.",
    "Therefore, the technological gap identified in this study was the limited availability of a context-specific system that integrates academic and personal characteristics to generate ranked and explainable university degree programme recommendations for Kenyan form-four leavers. This study addressed the gap by developing a Random Forest-based decision support system that uses students' academic performance, interests, skills, strengths, and career aspirations to generate three ranked programme recommendations together with SHAP-based recommendation explanations and supporting programme information.",
)
R(
    "The conceptual framework illustrates how the proposed decision support system transforms student information into personalized university degree programme recommendations. The process begins when a form-four leaver provides personal and academic information, including KCSE performance, interests, skills, strengths, and career aspirations. These inputs are validated and prepared before being analysed by the recommendation model.",
    "The conceptual framework illustrates how the developed decision support system transforms student information into personalized university degree programme recommendations. The process begins when a form-four leaver provides personal and academic information, including KCSE performance, interests, skills, strengths, and career aspirations. These inputs are validated and prepared before being analysed by the recommendation model.",
)

# ---------------------------------------------------------------------------
# Chapter 3: Methodology
# ---------------------------------------------------------------------------
R(
    "This chapter presents the methodology that will be used in the development and evaluation of the proposed Random Forest-based university degree programme recommendation system. It describes the research paradigm, data acquisition and processing procedures, model training and evaluation, and the software development methodology. The chapter also presents the system analysis and design approaches, expected system deliverables, and the tools and techniques that will be used to develop and integrate the recommendation system.",
    "This chapter presents the methodology that was used in the development and evaluation of the Random Forest-based university degree programme recommendation system. It describes the research paradigm, data acquisition and processing procedures, model training and evaluation, and the software development methodology. The chapter also presents the system analysis and design approaches, the system deliverables, and the tools and techniques used to develop and integrate the recommendation system.",
)
R(
    "This study will adopt an experimental research paradigm to guide the development and evaluation of the proposed Random Forest-based university degree programme recommendation system. The experimental approach will allow the researcher to develop and test a Random Forest-based recommendation model using student profile data. The model will be evaluated to determine its effectiveness in generating suitable university degree programme recommendations based on students' academic performance, interests, skills, strengths, and career aspirations. The study will use synthetic student profiles because a suitable labelled dataset linking these characteristics to university degree programmes is not readily available. The Cross-Industry Standard Process for Data Mining (CRISP-DM) will be used to guide the machine learning component through the stages of business understanding, data understanding, data preparation, modelling, evaluation, and deployment (Shimaoka et al., 2024). The modelling stage will focus specifically on the Random Forest algorithm, while the evaluation stage will assess its performance using accuracy, precision, recall, F1-score, and a confusion matrix. The experimental approach will provide a structured basis for determining whether the developed recommendation model can generate relevant and explainable university degree programme recommendations for Kenyan form-four leavers.",
    "This study adopted an experimental research paradigm to guide the development and evaluation of the Random Forest-based university degree programme recommendation system. The experimental approach allowed the researcher to develop and test a Random Forest-based recommendation model using student profile data. The model was evaluated to determine its effectiveness in generating suitable university degree programme recommendations based on students' academic performance, interests, skills, strengths, and career aspirations. The study used synthetic student profiles because a suitable labelled dataset linking these characteristics to university degree programmes is not readily available. The Cross-Industry Standard Process for Data Mining (CRISP-DM) was used to guide the machine learning component through the stages of business understanding, data understanding, data preparation, modelling, evaluation, and deployment (Shimaoka et al., 2024). The modelling stage focused specifically on the Random Forest algorithm, while the evaluation stage assessed its performance using accuracy, precision, recall, and F1-score, as reported in Chapter 5. The experimental approach provided a structured basis for determining whether the developed recommendation model could generate relevant and explainable university degree programme recommendations for Kenyan form-four leavers.",
)
R(
    "Data acquisition will primarily involve the generation of synthetic student profiles using a Large Language Model (LLM). The synthetic profiles will be developed using information from trusted educational and programme-related sources, including KUCCPS programme and admission information, Commission for University Education (CUE) information, university programme descriptions, and recognised career guidance resources. The profiles will represent realistic Kenyan form-four leaver characteristics, including KCSE performance, interests, skills, strengths, preferred subjects, and career aspirations.",
    "Data acquisition primarily involved the generation of synthetic student profiles using a deterministic, rule-based generator informed by trusted educational and programme-related sources, including KUCCPS programme and admission information, Commission for University Education (CUE) information, university programme descriptions, and recognised career guidance resources. The profiles represent realistic Kenyan form-four leaver characteristics, including KCSE performance across 7 to 9 subjects, interests, skills, strengths, preferred subjects, and career aspirations, as detailed in Chapter 5.",
)
R(
    "Each profile will be assigned a university degree programme label using predefined programme suitability criteria derived from the academic requirements and characteristics of the selected degree programmes. The selected university degree programmes will represent the classification classes used by the Random Forest model. Academic eligibility will first be considered using relevant KCSE subject and grade requirements. For programmes for which the student meets the academic requirements, the student's interests, skills, strengths, and career aspirations will then be matched against the characteristics associated with the programme. The programme with the strongest overall suitability based on the predefined criteria will be assigned as the profile label. The generated profiles and assigned labels will then be reviewed against the source information to identify inconsistencies and improve dataset quality. The LLM will therefore be used primarily for generating realistic student profiles rather than independently determining the programme labels.",
    "Each profile was assigned a university degree programme label using predefined programme suitability criteria derived from the academic requirements and characteristics of the 11 catalogued degree programmes, which represent the classification classes used by the Random Forest model. Academic eligibility was first considered using relevant KCSE subject and grade requirements. For programmes for which the student met the academic requirements, the student's interests, skills, strengths, and career aspirations were then matched against the characteristics associated with the programme. The programme with the strongest overall suitability based on the predefined criteria was assigned as the profile label, with ties broken by the closest KUCCPS Cluster Weighted Points fit, as detailed in Chapter 5.",
)
R(
    "Following data acquisition, the generated dataset will undergo preprocessing to improve its quality, consistency, and suitability for machine learning. The preprocessing stage will involve identifying and handling missing values, removing duplicate records, correcting inconsistent entries, and standardizing categorical values. Categorical attributes such as interests, skills, strengths, and career aspirations will be encoded into numerical representations suitable for the Random Forest model, while KCSE performance and other structured academic attributes will be represented using appropriate numerical or ordinal values. Relevant features will then be selected based on their contribution to programme recommendation. The processed student profiles will be organized into input features and university degree programme labels. The prepared dataset will then be divided into training and testing sets using a stratified split to maintain the distribution of programme classes across both sets. Duplicate or highly similar synthetic profiles will be identified and removed before splitting to reduce the risk of information overlap between the training and testing data. The testing set will remain completely unseen during model training, feature selection, preprocessing parameter estimation, and hyperparameter tuning. Cross-validation and model tuning will be performed using only the training data, while the held-out testing set will be used once for final performance evaluation.",
    "Following data acquisition, the generated dataset underwent preprocessing to improve its quality, consistency, and suitability for machine learning. Categorical attributes such as interests, skills, strengths, and career aspirations were one-hot encoded into numerical representations suitable for the Random Forest model, while KCSE performance and other structured academic attributes were represented as ordinal grade points. The processed student profiles were organized into input features and university degree programme labels, and the prepared dataset was divided into training and testing sets using an 80/20 stratified split to maintain the distribution of programme classes across both sets. The testing set remained completely unseen during model training, feature selection, and hyperparameter tuning. Cross-validation and model tuning were performed using only the training data, while the held-out testing set was used once for the final performance evaluation reported in Chapter 5.",
)
R(
    "The processed dataset will be used to train a Random Forest classification model for university degree programme recommendation. The dataset will be divided into training and testing subsets using a stratified split to maintain representation of the programme classes. The training data will be used to learn relationships between student characteristics and the assigned university degree programme labels. The model will use selected features including KCSE performance, interests, skills, strengths, and career aspirations to identify patterns associated with suitable degree programmes.",
    "The processed dataset was used to train a Random Forest classification model for university degree programme recommendation. The training data was used to learn relationships between student characteristics and the assigned university degree programme labels. The model used selected features, including KCSE performance across nine cluster-relevant subjects, interests, skills, strengths, and career aspirations, to identify patterns associated with suitable degree programmes.",
)
R(
    "Random Forest will be selected because its ensemble of decision trees can model non-linear relationships between multiple student characteristics and is suitable for structured datasets containing both numerical and encoded categorical features (Diamantopoulou et al., 2025). Relevant hyperparameters, such as the number of trees and maximum tree depth, will be tuned using cross-validation on the training data to improve generalization and reduce overfitting, while the testing data will remain untouched until final evaluation. The final trained model will then be evaluated using previously unseen test data before being integrated into the recommendation component of the system.",
    "Random Forest was selected because its ensemble of decision trees can model non-linear relationships between multiple student characteristics and is suitable for structured datasets containing both numerical and encoded categorical features (Diamantopoulou et al., 2025). Relevant hyperparameters, including the number of trees, maximum tree depth, and minimum samples per leaf, were tuned by grid search with cross-validation on the training data to improve generalization and reduce overfitting, while the testing data remained untouched until final evaluation. The final trained model was then evaluated using previously unseen test data, as reported in Chapter 5, before being integrated into the recommendation component of the system.",
)
R(
    "The trained Random Forest model will be validated and tested using the testing dataset that was not used during model training or hyperparameter tuning. The evaluation will determine how effectively the model can classify student profiles into suitable university degree programmes and how well it generalizes to previously unseen data. The model will be evaluated using accuracy, precision, recall, and F1-score to provide a balanced assessment of its classification performance (Rainio et al., 2024). A confusion matrix will also be used to examine the distribution of correct and incorrect predictions across the different programme classes and identify classes that may be difficult for the model to distinguish. The Random Forest model will produce prediction probabilities for the programme classes, which will be ranked in descending order to select the three highest-ranked recommendations for each student. SHAP will be used to generate explanations by identifying the student-profile features that contributed most to individual programme predictions. The explanations will be reviewed for clarity and relevance to the student's profile, while the final model performance will be reported using the unseen testing dataset.",
    "The trained Random Forest model was validated and tested using the testing dataset that was not used during model training or hyperparameter tuning. The model was evaluated using accuracy, precision, recall, and F1-score to provide a balanced assessment of its classification performance (Rainio et al., 2024), reported per programme class in Chapter 5 to identify classes the model found difficult to distinguish. At inference time, the Random Forest model produces prediction probabilities for each of the 11 programme classes, ranked in descending order to select the three highest-ranked recommendations for each student. SHAP is used to generate explanations by identifying the student-profile features that contributed most to each individual programme prediction, which are surfaced to the student in the mobile application together with the recommendation.",
)
R(
    "The Scrum-for-One Agile methodology will be used to guide the development of the proposed system. Scrum-for-One adapts the Scrum framework for an individual developer while maintaining its iterative and incremental development principles (Towhidnejad & Hagerty-Posell, 2025). This approach will allow the system to be developed in manageable stages, with each sprint focusing on a specific set of system requirements and functionalities. The development will involve activities such as backlog creation, sprint planning, implementation, testing, sprint review, and incorporation of improvements into subsequent sprints. This approach is suitable for the project because the system consists of several components, including student registration and profile management, data processing, the Random Forest recommendation model, recommendation explanations, programme information, and feedback functionality. CRISP-DM will be used specifically to guide the machine-learning component, while Scrum-for-One will guide the development and integration of the complete software system.",
    "The Scrum-for-One Agile methodology was used to guide the development of the system. Scrum-for-One adapts the Scrum framework for an individual developer while maintaining its iterative and incremental development principles (Towhidnejad & Hagerty-Posell, 2025). This approach allowed the system to be developed in manageable stages, with each sprint focusing on a specific set of system requirements and functionalities, involving backlog creation, sprint planning, implementation, testing, sprint review, and incorporation of improvements into subsequent sprints. This approach was suitable for the project because the system consists of several components, including student registration and profile management, data processing, the Random Forest recommendation model, recommendation explanations, programme information, and feedback functionality. CRISP-DM was used specifically to guide the machine-learning component, while Scrum-for-One guided the development and integration of the complete software system.",
)
R(
    "Scrum-for-One was selected because it provides a structured yet flexible approach for developing the system as an individual developer. The methodology supports iterative development, continuous testing, regular review, and progressive refinement of requirements and functionality. Each sprint will produce a functional increment of the system, allowing problems to be identified and addressed before proceeding to subsequent development activities. This is particularly suitable for the project because the recommendation model and other system components may require refinement during development and testing. The use of Scrum-for-One also provides a clear way of organizing development tasks through a product backlog and sprint backlog while maintaining the benefits of an Agile development approach. Therefore, Scrum-for-One will guide the software development process, while CRISP-DM will provide the structured process for developing and evaluating the Random Forest recommendation model. ",
    "Scrum-for-One was selected because it provided a structured yet flexible approach for developing the system as an individual developer. The methodology supported iterative development, continuous testing, regular review, and progressive refinement of requirements and functionality. Each sprint produced a functional increment of the system, allowing problems to be identified and addressed before proceeding to subsequent development activities, which was particularly useful because the recommendation model and other system components required refinement during development and testing. Scrum-for-One guided the software development process, while CRISP-DM provided the structured process for developing and evaluating the Random Forest recommendation model.",
)
R(
    "The product backlog will contain all identified requirements, features, and tasks required to develop the proposed university degree programme recommendation system. These requirements will be derived from the research objectives, literature review, existing systems, and identified user needs. The backlog will include core features such as user authentication, student profile management, student assessment, data preprocessing, Random Forest model development, programme recommendation, recommendation explanations, programme information, feedback functionality, system testing, and documentation. The backlog will be continuously reviewed, refined, and prioritized according to system importance and development progress.",
    "The product backlog contained all identified requirements, features, and tasks required to develop the university degree programme recommendation system, derived from the research objectives, literature review, existing systems, and identified user needs. The backlog included core features such as user authentication, student profile management, student assessment, data preprocessing, Random Forest model development, programme recommendation, recommendation explanations, programme information, feedback functionality, system testing, and documentation, and was continuously reviewed, refined, and prioritized according to system importance and development progress.",
)
R(
    "Sprint planning will involve selecting and prioritizing items from the product backlog for development within a specific sprint. Each sprint will have a defined goal and time frame, with selected requirements broken down into smaller and manageable tasks. The planning process will consider the dependencies between activities, such as completing student profile and data-processing functionality before integrating the Random Forest recommendation model. This will ensure that each sprint has clear objectives and produces an achievable functional outcome.",
    "Sprint planning involved selecting and prioritizing items from the product backlog for development within a specific sprint. Each sprint had a defined goal and time frame, with selected requirements broken down into smaller and manageable tasks. The planning process considered the dependencies between activities, such as completing student profile and data-processing functionality before integrating the Random Forest recommendation model, ensuring that each sprint had clear objectives and produced an achievable functional outcome.",
)
R(
    "The sprint backlog will contain the specific tasks selected from the product backlog for completion during a particular sprint. These tasks will define the functionality to be developed, tested, and integrated within that sprint. Depending on the sprint goal, tasks may include implementing authentication, developing the student profile, preparing data, training the Random Forest model, implementing recommendation generation, developing the explanation and programme information components, or conducting system testing. The sprint backlog will be monitored and updated throughout the sprint as tasks are completed or adjusted based on development progress.",
    "The sprint backlog contained the specific tasks selected from the product backlog for completion during a particular sprint, including implementing authentication, developing the student profile, preparing data, training the Random Forest model, implementing recommendation generation, developing the explanation and programme information components, and conducting system testing. The sprint backlog was monitored and updated throughout each sprint as tasks were completed or adjusted based on development progress.",
)
R(
    "The system will be developed through short, iterative sprints, with each sprint focusing on a specific set of prioritized requirements. Development activities within a sprint will include analysis, design, implementation, integration, and testing of the selected functionality. The sprints will progressively develop the system, beginning with foundational features such as authentication and student profiling, followed by data processing and Random Forest recommendation functionality, and then supporting features such as recommendation explanations, programme information, feedback, and system integration. Each sprint will produce a functional increment that can be reviewed before proceeding to subsequent development activities.",
    "The system was developed through short, iterative sprints, with each sprint focusing on a specific set of prioritized requirements, including analysis, design, implementation, integration, and testing of the selected functionality. The sprints progressively developed the system, beginning with foundational features such as authentication and student profiling, followed by data processing and Random Forest recommendation functionality, and then supporting features such as recommendation explanations, programme information, feedback, and system integration. Each sprint produced a functional increment that was reviewed before proceeding to subsequent development activities.",
)
R(
    "At the end of each sprint, a Sprint Review will be conducted to assess the completed functionality against the sprint objectives and identified requirements. The implemented features will be tested to determine whether they function as expected and whether they meet the intended system requirements. Any errors, incomplete functionality, or areas requiring improvement will be documented and considered when refining the product backlog and planning subsequent sprints. For the recommendation component, the review will also consider whether the implemented functionality correctly processes student profiles and produces the intended programme recommendations and explanations.",
    "At the end of each sprint, a Sprint Review was conducted to assess the completed functionality against the sprint objectives and identified requirements. The implemented features were tested to determine whether they functioned as expected and met the intended system requirements. Any errors, incomplete functionality, or areas requiring improvement were documented and considered when refining the product backlog and planning subsequent sprints. For the recommendation component, the review also considered whether the implemented functionality correctly processed student profiles and produced the intended programme recommendations and explanations.",
)
R(
    "An increment will be the functional and tested output produced at the end of each sprint. Each increment will add new or improved functionality to the existing system while maintaining the functionality developed in previous sprints. For example, an increment may initially provide authentication and student profiling, while later increments may add data processing, Random Forest-based recommendations, recommendation explanations, programme information, and feedback functionality. The increments will be progressively integrated and tested until the complete system meets the defined functional and project requirements.",
    "An increment was the functional and tested output produced at the end of each sprint, each adding new or improved functionality to the existing system while maintaining the functionality developed in previous sprints. The initial increment provided authentication and student profiling, while later increments added data processing, Random Forest-based recommendations, recommendation explanations, programme information, and feedback functionality. The increments were progressively integrated and tested until the complete system met the defined functional and project requirements.",
)
R(
    "The System Analysis and Design phase will define the structure, functionality, data relationships, and overall architecture of the proposed university degree programme recommendation system. This phase will translate the identified system requirements into models and designs that will guide implementation. The analysis and design will focus on student interactions, administrator activities, student profile data, recommendation generation, programme information, data storage, and communication between the application and the Random Forest recommendation model. ",
    "The System Analysis and Design phase defined the structure, functionality, data relationships, and overall architecture of the university degree programme recommendation system, translating the identified system requirements into models and designs that guided implementation. The analysis and design focused on student interactions, administrator activities, student profile data, recommendation generation, programme information, data storage, and communication between the application and the Random Forest recommendation model.",
)
R(
    "The Use Case Diagram will illustrate the interactions between the system and its primary actors, the Student and System Administrator. The student will be able to register and log in, create and update a profile, provide academic and personal information, request programme recommendations, view the ranked recommendations and their explanations, access supporting programme information, and provide feedback on the recommendations and update their profile to request a new recommendation if the recommendations are unsuitable. The System Administrator will maintain programme and career information stored in the knowledge base to ensure that the supporting information presented to students remains relevant and up to date.",
    "The Use Case Diagram illustrates the interactions between the system and its primary actors, the Student and System Administrator. The student can register and log in, create and update a profile, provide academic and personal information, request programme recommendations, view the ranked recommendations and their explanations, access supporting programme information, and provide feedback on the recommendations and update their profile to request a new recommendation if the recommendations are unsuitable. The System Administrator maintains programme and career information stored in the knowledge base to ensure that the supporting information presented to students remains relevant and up to date.",
)
R(
    "The Sequence Diagram will illustrate the sequence of interactions when a student requests programme recommendation. It will show the student submitting profile information through the system, the data being processed and sent to the Random Forest model, the generation of ranked recommendations, retrieval of supporting programme information, and presentation of the results to the student.",
    "The Sequence Diagram illustrates the sequence of interactions when a student requests programme recommendation. It shows the student submitting profile information through the system, the data being processed and sent to the Random Forest model, the generation of ranked recommendations, retrieval of supporting programme information, and presentation of the results to the student.",
)
R(
    "The database schema will define how student profiles, degree programmes, programme information, recommendations, feedback, and administrator information will be stored and related. The schema will support the storage and retrieval of the information required for generating recommendations and presenting relevant programme details to students.",
    "The database schema defines how student profiles, degree programmes, programme information, recommendations, feedback, and administrator information are stored and related within Cloud Firestore, supporting the storage and retrieval of the information required for generating recommendations and presenting relevant programme details to students.",
)
R(
    "The system architecture will illustrate the interaction between the user interface, application layer, Random Forest recommendation model, and database. It will show how student profile information is submitted and processed, how the model generates the ranked recommendations, and how programme information is retrieved and presented alongside the recommendations.",
    "The system architecture illustrates the interaction between the user interface, application layer, Random Forest recommendation model, and database. It shows how student profile information is submitted and processed, how the model generates the ranked recommendations, and how programme information is retrieved and presented alongside the recommendations.",
)
R(
    "Wireframes will be developed to represent the main interfaces before implementation and establish the intended user flow. These will include the registration and login pages, student profile page, recommendation request and results pages, programme information page, and feedback interface. The wireframes will help ensure that the main functions are clearly organized and easy for students to navigate.",
    "Wireframes were developed to represent the main interfaces before implementation and establish the intended user flow, including the registration and login pages, student profile page, recommendation request and results pages, programme information page, and feedback interface, helping ensure that the main functions are clearly organized and easy for students to navigate.",
)
R(
    "Comprehensive system documentation will be developed to describe the requirements, design, implementation, testing, and operation of the proposed university degree programme recommendation system. The documentation will include system architecture, database design, model development and evaluation procedures, user instructions, and maintenance information. It will provide a reference for understanding, operating, testing, and maintaining the developed system.",
    "Comprehensive system documentation was developed to describe the requirements, design, implementation, testing, and operation of the university degree programme recommendation system, including system architecture, database design, model development and evaluation procedures, user instructions, and maintenance information, providing a reference for understanding, operating, testing, and maintaining the developed system.",
)
R(
    "The Authentication Module will provide secure access to the system by allowing students and administrators to register, log in, and manage their accounts. It will implement authentication and authorization mechanisms to ensure that users can only access functions appropriate to their roles. The module will also support secure account and password management to improve system security and usability.",
    "The Authentication Module provides secure access to the system by allowing students and administrators to register, log in, and manage their accounts, using Firebase Authentication. It implements authorization mechanisms to ensure that users can only access functions appropriate to their roles, and supports secure account and password management to improve system security and usability.",
)
R(
    "The Student Assessment and Profiling Module will collect and manage the information required to create a student profile for programme recommendation. The information will include KCSE performance, interests, skills, strengths, and career aspirations. Students will be able to enter and update their profile information, which will then be processed into the features required by the Random Forest recommendation model.",
    "The Student Assessment and Profiling Module collects and manages the information required to create a student profile for programme recommendation, including KCSE performance across 7 to 9 subjects, interests, skills, strengths, and career aspirations. Students enter and update their profile information, which is then processed into the features required by the Random Forest recommendation model.",
)
R(
    "The Recommendation Module will serve as the core machine learning component of the system. It will use the trained Random Forest classification model to analyse the student's profile and generate prediction probabilities for the selected university degree programmes. The probabilities will then be ranked to identify the three highest-ranked programmes for presentation to the student. SHAP will be used to generate explanations showing the main student-profile features that contributed to each recommendation. The module will therefore provide both ranked recommendations and model-based explanations to support informed decision-making.",
    "The Recommendation Module serves as the core machine learning component of the system. It uses the trained Random Forest classification model to analyse the student's profile and generate prediction probabilities for each of the 11 catalogued university degree programmes, which are ranked to identify the three highest-ranked programmes for presentation to the student. SHAP generates explanations showing the main student-profile features that contributed to each recommendation, so the module provides both ranked recommendations and model-based explanations to support informed decision-making.",
)
R(
    "The Programme Information and Support Module will provide additional information about the recommended university degree programmes to support students in evaluating their options. The information will include programme entry requirements, universities offering the programme, related career pathways, recommended skills, certifications, and relevant learning resources. This information will be retrieved from the programme knowledge base and presented alongside the recommendations, while recommendation explanations will be generated separately from the Random Forest model using SHAP. Students will be able to review the recommendations, explanations, and supporting programme information before providing feedback. If the student considers the recommendations unsuitable, they will be able to update their profile information and request a new recommendation based on the revised inputs. The module will support informed decision-making without functioning as a separate academic learning or tutoring system.",
    "The Programme Information and Support Module provides additional information about the recommended university degree programmes to support students in evaluating their options, including programme entry requirements, universities offering the programme, related career pathways, recommended skills, certifications, and relevant learning resources. This information is retrieved from the programme knowledge base and presented alongside the recommendations, while recommendation explanations are generated separately from the Random Forest model using SHAP. Students can review the recommendations, explanations, and supporting programme information before providing feedback, and if a student considers the recommendations unsuitable, they can update their profile information and request a new recommendation based on the revised inputs. The module supports informed decision-making without functioning as a separate academic learning or tutoring system.",
)
R(
    "Python will be used to develop the machine learning component of the proposed system. It will support data processing, preparation of student profile data, training of the Random Forest classification model, and generation of university degree programme recommendations. Python is suitable for the project because it provides extensive libraries and tools for machine learning and data analysis.",
    "Python was used to develop the machine learning component and the FastAPI backend of the system, supporting data processing, preparation of student profile data, training of the Random Forest classification model, and generation of university degree programme recommendations. Python was suitable for the project because it provides extensive libraries and tools for machine learning, data analysis, and web API development.",
)
R(
    "Scikit-Learn will be used to build, train, and evaluate the machine learning model. It provides a range of reliable algorithms for classification, regression, and clustering, which are useful for recommendation systems. It also includes tools for data preprocessing, feature selection, and model evaluation, making the development process more efficient. In addition, it supports 19 evaluation metrics such as accuracy, precision, recall, and F1-score to assess model performance. Its simplicity and strong documentation make it suitable for rapid development and testing of the proposed system.",
    "Scikit-Learn was used to build, train, and evaluate the Random Forest model, including its OneHotEncoder for categorical features, GridSearchCV for hyperparameter tuning, and evaluation metrics such as accuracy, precision, recall, and F1-score. Its simplicity and strong documentation made it suitable for rapid development and testing of the recommendation model.",
)
R(
    "SHAP will be used to generate explanations for individual Random Forest predictions. It will identify the contribution of relevant student-profile features to the predicted university degree programmes, allowing the system to present understandable reasons for the recommendations. This will support transparency by showing students which aspects of their profiles contributed most to each recommendation.",
    "SHAP is used to generate explanations for individual Random Forest predictions, identifying the contribution of relevant student-profile features to the predicted university degree programmes and allowing the system to present understandable reasons for the recommendations. This supports transparency by showing students which aspects of their profiles contributed most to each recommendation.",
)
R(
    "FastAPI will be used to develop the backend API that connects the Flutter application with the Random Forest recommendation model and other backend services. It will receive student profile information from the application, pass the processed information to the recommendation model, and return the generated programme recommendations and relevant information to the application.",
    "FastAPI was used to develop the backend REST API that connects the Flutter application with the Random Forest recommendation model, Firebase Authentication, and Cloud Firestore. It receives student profile information from the application, verifies the caller's identity, passes the processed information to the recommendation model, and returns the generated programme recommendations and relevant information to the application.",
)
R(
    "Flutter will be used to develop the mobile application interface of the proposed system. It will provide the interfaces through which students can register and log in, create and manage their profiles, submit their academic and personal information, request recommendations, view ranked programme recommendations, access programme information, and provide feedback.",
    "Flutter was used to develop the mobile application interface of the system, providing the interfaces through which students register and log in, create and manage their profiles, submit their academic and personal information, request recommendations, view ranked programme recommendations, access programme information, and provide feedback.",
)
R(
    "Firebase will be used to provide authentication and data storage services for the system. It will support user authentication and storage of student profiles, degree programme information, recommendations, and student feedback. The stored feedback will support evaluation of recommendation relevance and will allow the system to record whether students consider the recommendations suitable. The feedback will not be used for automatic retraining of the Random Forest model during the study. Firebase will enable the application to securely store and retrieve the information required by the system.",
    "Firebase provides authentication and data storage services for the system, through Firebase Authentication and Cloud Firestore. It supports user authentication and storage of student profiles, degree programme information, recommendations, and student feedback. The stored feedback supports evaluation of recommendation relevance and allows the system to record whether students consider the recommendations suitable; the feedback was not used for automatic retraining of the Random Forest model during this study. Firebase enables the application to securely store and retrieve the information required by the system, with Firestore security rules restricting direct client access so that all reads and writes are mediated by the authenticated FastAPI backend.",
)
R(
    "Visual Studio Code will be used as the primary development environment for implementing, testing, and debugging the proposed system. It will support development of the Python machine learning component, FastAPI backend, and Flutter application through its development tools and extensions.",
    "Visual Studio Code was used as the primary development environment for implementing, testing, and debugging the system, supporting development of the Python machine learning component, FastAPI backend, and Flutter application through its development tools and extensions.",
)

# ---------------------------------------------------------------------------
# Chapter 4: System Analysis and Design
# ---------------------------------------------------------------------------
R(
    "This chapter details the analysis and design phase of the UniGuide University Degree Programme Recommendation System, transforming the project's objectives into a concrete technical plan. It begins with a comprehensive requirement analysis that outlines the needs of the main stakeholders and defines the system's functional and non-functional requirements. Following the requirements, this chapter presents a series of system analysis and design diagrams that provide a logical and visual representation of the proposed system and guide its implementation.",
    "This chapter details the analysis and design phase of the UniGuide University Degree Programme Recommendation System, transforming the project's objectives into a concrete technical plan that guided implementation. It begins with a comprehensive requirement analysis that outlines the needs of the main stakeholders and defines the system's functional and non-functional requirements. Following the requirements, this chapter presents a series of system analysis and design diagrams that provide a logical and visual representation of the system.",
)
R(
    "The system analysis section focuses on understanding what the system should do and how users interact with it. The system design section then presents the proposed technical architecture, data organisation, and user interface through the relevant design diagrams and wireframes.",
    "The system analysis section focuses on understanding what the system should do and how users interact with it. The system design section then presents the technical architecture, data organisation, and user interface through the relevant design diagrams and wireframes.",
)
R(
    "The requirement analysis was conducted with consideration of the needs of the primary users of UniGuide, particularly students who require support when selecting suitable university degree programmes. The system is intended to address the challenge faced by form-four leavers who may have difficulty identifying degree programmes that align with their academic performance, personal interests, skills, strengths, and career aspirations.",
    "The requirement analysis was conducted with consideration of the needs of the primary users of UniGuide, particularly students who require support when selecting suitable university degree programmes. The system addresses the challenge faced by form-four leavers who may have difficulty identifying degree programmes that align with their academic performance, personal interests, skills, strengths, and career aspirations.",
)
R(
    "Students currently consider different sources of information when making programme choices. These include their KCSE performance, programme admission requirements, personal interests, skills, and future career goals. UniGuide aims to bring these factors together within one platform and provide personalised degree programme recommendations supported by understandable explanations and relevant programme information.",
    "Students currently consider different sources of information when making programme choices, including their KCSE performance, programme admission requirements, personal interests, skills, and future career goals. UniGuide brings these factors together within one platform and provides personalised degree programme recommendations supported by understandable explanations and relevant programme information.",
)
R(
    "The main stakeholders of the proposed system are students and administrators. Students use the system to provide their profile information, request recommendations, view recommended degree programmes and their explanations, access programme information, and provide feedback. Administrators are responsible for maintaining the programme-related information used by the system, including admission requirements and career pathway information.",
    "The main stakeholders of the system are students and administrators. Students use the system to provide their profile information, request recommendations, view recommended degree programmes and their explanations, access programme information, and provide feedback. Administrators are responsible for maintaining the programme-related information used by the system, including admission requirements, minimum grades, and programme descriptions.",
)
R(
    "FR-01: Authentication Module-This provides a way of identifying users and controlling their access to the system. First-time users will be required to register before accessing the system, while returning users will log in using their registered credentials. The system will support separate access privileges for students and administrators, ensuring that administrative functions are only available to authorised users. Firebase Authentication will be used to manage user authentication and account access.",
    "FR-01: Authentication Module-This provides a way of identifying users and controlling their access to the system. First-time users register before accessing the system, while returning users log in using their registered credentials. The system supports separate access privileges for students and administrators, ensuring that administrative functions are only available to authorised users. Firebase Authentication is used to manage user authentication and account access, with the administrator role resolved server-side rather than inferred from the account's email address.",
)
R(
    "FR-02: Student Profile Management-This will provide students with a way of entering and updating the information required for recommendations. The profile will include KCSE academic performance, interests, skills, strengths, and career aspirations. The information provided will be used as the input for generating personalised degree programme recommendations.",
    "FR-02: Student Profile Management-This provides students with a way of entering and updating the information required for recommendations. The profile includes KCSE academic performance across 7 to 9 subjects, interests, skills, strengths, and career aspirations. The information provided is used as the input for generating personalised degree programme recommendations.",
)
R(
    "FR-03: Degree Programme Recommendation- This is the main intelligence of the system. It will analyse the student's academic and personal profile using the trained Random Forest model to identify suitable university degree programmes. The predicted programmes will be ranked according to their confidence scores, and the highest-ranked programmes will be presented to the student.",
    "FR-03: Degree Programme Recommendation- This is the main intelligence of the system. It analyses the student's academic and personal profile using the trained Random Forest model to identify suitable university degree programmes. The predicted programmes are ranked according to their confidence scores, and the three highest-ranked programmes are presented to the student, together with a SHAP-based explanation for each.",
)
R(
    "FR-04: Feedback and Programme Management- This will allow students to provide feedback on the relevance of their recommendations, while authorised administrators will manage programme information and admission requirements. Student feedback will support evaluation and future system improvement but will not automatically retrain the Random Forest model within this study.",
    "FR-04: Feedback and Programme Management- This allows students to provide feedback on the relevance of their recommendations, while authorised administrators manage programme information and admission requirements. Student feedback supports evaluation and future system improvement but does not automatically retrain the Random Forest model within this study.",
)
R(
    "FR-05: Recommendation Explanation and Programme Information-This will provide students with explanations and supporting information about their recommendations. SHAP will identify important student-profile features that contributed to individual predictions, while programme information will provide details such as admission requirements, university options, and potential career pathways.",
    "FR-05: Recommendation Explanation and Programme Information-This provides students with explanations and supporting information about their recommendations. SHAP identifies the student-profile features that contributed most to each prediction, while programme information provides details such as admission requirements, university options, and potential career pathways.",
)
R(
    "NFR-01: Security-The system handles personal student information and therefore requires secure authentication and access control. Firebase Authentication will be used to manage user authentication while access to administrative functions will be restricted to authorised administrators. ",
    "NFR-01: Security-The system handles personal student information and therefore requires secure authentication and access control. Firebase Authentication manages user authentication, access to administrative functions is restricted to authorised administrators, and Cloud Firestore security rules deny direct client access so that all data access is mediated by the authenticated backend.",
)
R(
    "NFR-02: Usability-the system should provide a simple and intuitive mobile interface. Students should be able to enter their information, request recommendations, understand the results, and access supporting programme information without unnecessary complexity. ",
    "NFR-02: Usability-the system provides a simple and intuitive mobile interface. Students can enter their information, request recommendations, understand the results, and access supporting programme information without unnecessary complexity.",
)
R(
    "NFR-03: Performance-The system should respond efficiently to student requests. Recommendation requests should be processed within a reasonable period while communication between the mobile application, FastAPI backend, machine-learning component, and Firebase remains efficient. ",
    "NFR-03: Performance-The system responds efficiently to student requests, with recommendation requests processed within a reasonable period while communication between the mobile application, FastAPI backend, machine-learning component, and Firebase remains efficient.",
)
R(
    "NFR-04: Reliability-The system should operate consistently and provide reliable recommendation results and programme information. Appropriate validation and error handling should be incorporated to reduce failures during profile submission, recommendation requests, and data retrieval. ",
    "NFR-04: Reliability-The system operates consistently and provides reliable recommendation results and programme information. Validation and error handling are incorporated throughout the mobile application and backend to reduce failures during profile submission, recommendation requests, and data retrieval, with clear error states shown to the user rather than the application silently failing.",
)
R(
    "NFR-05: Scalability-The system architecture should allow the application to accommodate an increasing number of student profiles, recommendations, programme records, and feedback records without requiring major changes to the overall architecture. ",
    "NFR-05: Scalability-The system architecture allows the application to accommodate an increasing number of student profiles, recommendations, programme records, and feedback records without requiring major changes to the overall architecture.",
)
R(
    "NFR-06: Compatibility-The proposed mobile application should operate consistently on supported mobile devices. The backend services should also maintain compatibility with the technologies used to access the system.",
    "NFR-06: Compatibility-The mobile application operates consistently on supported Android devices. The backend services also maintain compatibility with the technologies used to access the system.",
)
R(
    "IR-06: Secure Communication & Error Handling: The system shall enforce HTTPS encrypted communication between the front end and backend and return standardized, user-friendly error responses without exposing system infrastructure details.",
    "IR-06: Secure Communication & Error Handling: The system returns standardized, user-friendly error responses without exposing system infrastructure details. HTTPS encrypted communication between the front end and backend is required for any deployment beyond local development and testing, as discussed in Chapter 5.",
)
R(
    "This section presents the system analysis diagrams for UniGuide. The diagrams model the system's functional requirements, processes, interactions, and data relationships from a high-level perspective. They illustrate what the system does and how users interact with it.",
    "This section presents the system analysis diagrams for UniGuide. The diagrams model the system's functional requirements, processes, interactions, and data relationships from a high-level perspective, illustrating what the system does and how users interact with it.",
)
R(
    "The Use Case Diagram defines the functional scope and interaction boundaries for UniGuide. It addresses functional requirements FR-01 through FR-05 by outlining user interactions for both primary actors: Students and System Administrators.",
    "The Use Case Diagram defines the functional scope and interaction boundaries for UniGuide, addressing functional requirements FR-01 through FR-05 by outlining user interactions for both primary actors: Students and System Administrators.",
)
R(
    "The Activity Diagram details the operational workflow during a recommendation request. It addresses performance, error handling, and data processing requirements across the application layer, machine learning engine, and data storage.",
    "The Activity Diagram details the operational workflow during a recommendation request, addressing performance, error handling, and data processing requirements across the application layer, machine learning engine, and data storage.",
)
R(
    "The Sequence Diagram depicts the step-by-step message exchanges and temporal execution sequence between components. It addresses system integration requirements across the Flutter frontend, FastAPI backend, ML service, and Firestore database. ",
    "The Sequence Diagram depicts the step-by-step message exchanges and temporal execution sequence between components, addressing system integration requirements across the Flutter frontend, FastAPI backend, ML service, and Firestore database.",
)
R(
    "This section contains the system design diagrams that provide a more detailed technical representation of UniGuide. They illustrate how the system is structured internally and how the analysed requirements will be implemented.",
    "This section contains the system design diagrams that provide a more detailed technical representation of UniGuide, illustrating how the system is structured internally and how the analysed requirements were implemented.",
)
R(
    "The System Architecture Diagram outlines the high-level structural components and technical design of the platform. It addresses architectural requirements for system modularity, API communication, model processing, and data persistence.",
    "The System Architecture Diagram outlines the high-level structural components and technical design of the platform, addressing architectural requirements for system modularity, API communication, model processing, and data persistence.",
)
R(
    "The Database Schema diagram defines the NoSQL collection structure within Firebase Cloud Firestore. It addresses data storage requirements DR-01 through DR-06 for managing user accounts, profiles, course details, recommendations, explanations, and audit metadata.",
    "The Database Schema diagram defines the NoSQL collection structure within Firebase Cloud Firestore, addressing data storage requirements DR-01 through DR-06 for managing user accounts, profiles, course details, recommendations, explanations, and audit metadata.",
)
R(
    "To support the proposed system, the UI wireframes outline the layout, navigation flow, and core user interactions across the application. These low-fidelity schematics serve as a visual blueprint for the interface, ensuring an intuitive, accessible experience for the students.",
    "The UI wireframes outline the layout, navigation flow, and core user interactions across the application. These low-fidelity schematics served as a visual blueprint for the interface, and the implemented mobile application described in Chapter 5 follows this same navigation flow.",
)
R(
    "The wireframe designs establish clear structural progression that translates system requirements into a user-centred interface flow. The onboarding and profile screens streamline access and gather KCSE subject grades and preferences in a single step. The recommendation results screen replaces traditional black-box guidance with clear confidence scores, KUCCPS cluster cutoffs, and visual SHAP feature contribution bars that explain match drivers. Finally, the programme detail screens connect predictive matches to career pathways and offering universities, while the administrative dashboard and feedback tools support ongoing system evaluation and model maintenance.",
    "The wireframe designs established a clear structural progression that translated system requirements into a user-centred interface flow, which the implemented application preserved: the onboarding and profile screens gather KCSE subject grades and preferences, the recommendation results screen replaces traditional black-box guidance with clear confidence scores, KUCCPS cluster cutoffs, and visual SHAP feature contribution bars that explain match drivers, the programme detail screens connect predictive matches to career pathways and offering universities, and the administrative dashboard and feedback tools support ongoing system evaluation and catalogue maintenance.",
)

# These four are trailing sentences INSIDE much larger paragraphs (the
# KUCCPS/DIRAPATH/CareerExplorer/CareerVillage related-works descriptions),
# not standalone paragraphs, so they need substring replacement instead of a
# whole-paragraph match.
SUBSTRING_REPLACEMENTS = [
    (
        "These limitations highlight the need for a decision support system that integrates both academic performance and personal characteristics to generate personalized and explainable university degree programme recommendations, thereby supporting Kenyan form-four leavers in making informed programme choices before submitting their KUCCPS applications.",
        "These limitations highlighted the need for a decision support system that integrates both academic performance and personal characteristics to generate personalized and explainable university degree programme recommendations, which this study addressed to support Kenyan form-four leavers in making informed programme choices before submitting their KUCCPS applications.",
    ),
    (
        "These limitations demonstrate the need for the proposed Random Forest-based decision support system, which combines both academic and personal characteristics to recommend the three most suitable university degree programmes while providing explanations that support informed decision-making.",
        "These limitations demonstrated the need for the Random Forest-based decision support system developed in this study, which combines both academic and personal characteristics to recommend the three most suitable university degree programmes while providing explanations that support informed decision-making.",
    ),
    (
        "These limitations highlight the need for the proposed Random Forest-based decision support system, which generates personalized and explainable university degree programme recommendations by integrating academic performance with students' interests, skills, strengths, and career aspirations within the Kenyan context.",
        "These limitations highlighted the need for the Random Forest-based decision support system developed in this study, which generates personalized and explainable university degree programme recommendations by integrating academic performance with students' interests, skills, strengths, and career aspirations within the Kenyan context.",
    ),
    (
        "These limitations highlight the need for the proposed Random Forest-based decision support system, which uses machine learning to generate personalized, explainable, and ranked university degree programme recommendations tailored to Kenyan form-four leavers.",
        "These limitations highlighted the need for the Random Forest-based decision support system developed in this study, which uses machine learning to generate personalized, explainable, and ranked university degree programme recommendations tailored to Kenyan form-four leavers.",
    ),
]


def replace_substring_in_paragraph(paragraph, old, new):
    """Replaces `old` with `new` inside a paragraph's concatenated run text,
    writing the result into the first run and clearing the others, so
    formatting (font/size/bold) from the first run is preserved."""
    full_text = paragraph.text
    if old not in full_text:
        return False
    updated = full_text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(updated)
    return True


# ---------------------------------------------------------------------------
# Apply
# ---------------------------------------------------------------------------
doc = docx.Document(PATH)

applied = 0
for p in doc.paragraphs:
    key = p.text.strip()
    if key in REPLACEMENTS:
        new_text = REPLACEMENTS.pop(key)
        if p.runs:
            p.runs[0].text = new_text
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.add_run(new_text)
        applied += 1

substring_applied = 0
for p in doc.paragraphs:
    for old, new in SUBSTRING_REPLACEMENTS:
        if replace_substring_in_paragraph(p, old, new):
            substring_applied += 1

print(f"Applied {applied} whole-paragraph replacements and {substring_applied} substring replacements.")
if REPLACEMENTS:
    print(f"WARNING: {len(REPLACEMENTS)} whole-paragraph replacement(s) were never matched:")
    for k in REPLACEMENTS:
        print("  -", k[:90].replace("\n", " "))
if substring_applied < len(SUBSTRING_REPLACEMENTS):
    print(f"WARNING: only {substring_applied}/{len(SUBSTRING_REPLACEMENTS)} substring replacements matched.")

doc.save(PATH)
print("Saved:", PATH)
